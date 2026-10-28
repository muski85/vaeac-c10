"""Copy NRMS / AE results from results.csv into a copy of the course Table-NRMS-AE.xlsx.

python fill_table.py --table "Table-NRMS-AE.xlsx" --results results/results.csv --out Table-NRMS-AE_C10.xlsx
"""
import argparse

import openpyxl
import pandas as pd

# The table spells a few dataset names differently from the folder names.
ALIAS = {"Bupa": "BUPA", "TTTTEG": "TTTEG", "CREDIT": "Credit"}

ap = argparse.ArgumentParser()
ap.add_argument("--table", required=True)
ap.add_argument("--results", required=True)
ap.add_argument("--out", required=True)
args = ap.parse_args()

res = pd.read_csv(args.results)
rows = {}
for _, r in res.iterrows():
    ds, rest = r["file"].split("_", 1)
    rows[f"{ALIAS.get(ds, ds)}_{rest}"] = r

wb = openpyxl.load_workbook(args.table)
ws = wb.active
filled = 0
for i in range(2, ws.max_row + 1):
    r = rows.get(ws.cell(i, 1).value)
    if r is None:
        continue
    if pd.notna(r.get("NRMS")):
        ws.cell(i, 2).value = round(float(r["NRMS"]), 4)
    if pd.notna(r.get("AE")):
        ws.cell(i, 3).value = round(float(r["AE"]), 4)
    ws.cell(i, 4).value = r["key_parameters"]
    filled += 1
wb.save(args.out)
print(f"filled {filled} rows -> {args.out}")
